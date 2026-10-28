"""Exportação da transcrição: TXT, SRT, VTT, DOCX, JSON."""

from __future__ import annotations

import json
import math
from datetime import datetime
from pathlib import Path

from .engine import SENTENCE_END, Transcript, Utterance, Word, fmt_time

FORMATS = {
    "txt": "Texto (.txt)",
    "docx": "Word (.docx)",
    "srt": "Legenda SRT (.srt)",
    "vtt": "Legenda WebVTT (.vtt)",
    "json": "JSON com tempos por palavra (.json)",
}

LANG_NAMES = {"pt": "Português", "en": "Inglês"}


def export(tr: Transcript, fmt: str, dest: Path) -> Path:
    path = dest.with_suffix("." + fmt)
    {"txt": _txt, "srt": _srt, "vtt": _vtt, "docx": _docx, "json": _json}[fmt](tr, path)
    return path


def _prefix(tr: Transcript, u: Utterance) -> str:
    return f"{tr.speaker_label(u.speaker)}: " if tr.diarized else ""


def to_text(tr: Transcript) -> str:
    lines = []
    for u in tr.utterances:
        lines.append(f"[{fmt_time(u.start)}] {_prefix(tr, u)}{u.text}")
    return "\n\n".join(lines) + "\n"


def _txt(tr: Transcript, path: Path) -> None:
    path.write_text(to_text(tr), encoding="utf-8")


def _join(words: list[Word]) -> str:
    return "".join(w.text for w in words).strip()


def _split_balanced(words: list[Word], max_chars: int, max_dur: float) -> list[list[Word]]:
    """Divide uma frase longa em partes de tamanho parecido, preferindo cortar após vírgulas."""
    length = len(_join(words))
    dur = words[-1].end - words[0].start
    n = max(math.ceil(length / max_chars), math.ceil(dur / max_dur), 1)
    if n == 1 or len(words) < 2:
        return [words]
    target = length / n
    parts: list[list[Word]] = []
    cur: list[Word] = []
    for i, w in enumerate(words):
        cur.append(w)
        size = len(_join(cur))
        remaining = len(words) - i - 1
        if len(parts) < n - 1 and remaining and (
            size >= target or (size >= target * 0.6 and w.text.strip().endswith((",", ";", ":")))
        ):
            parts.append(cur)
            cur = []
    if cur:
        parts.append(cur)
    return parts


def _cues(tr: Transcript, max_chars: int = 80, max_dur: float = 7.0):
    """Quebra cada fala em legendas curtas, nunca misturando falantes."""
    for u in tr.utterances:
        sentences: list[list[Word]] = [[]]
        for w in u.words:
            sentences[-1].append(w)
            if w.text.strip().endswith(SENTENCE_END):
                sentences.append([])
        groups: list[list[Word]] = []
        for sent in filter(None, sentences):
            # Junta frases curtas na mesma legenda quando cabem.
            if groups and len(_join(groups[-1] + sent)) <= max_chars \
                    and sent[-1].end - groups[-1][0].start <= max_dur:
                groups[-1] += sent
            else:
                groups.extend(_split_balanced(sent, max_chars, max_dur))
        for g in groups:
            yield g[0].start, g[-1].end, u.speaker, _join(g)


def _wrap(text: str, width: int = 42) -> str:
    if len(text) <= width:
        return text
    mid = len(text) // 2
    left, right = text.rfind(" ", 0, mid + 1), text.find(" ", mid)
    cut = left if right == -1 or (left != -1 and mid - left <= right - mid) else right
    if cut <= 0:
        return text
    return text[:cut].strip() + "\n" + text[cut:].strip()


def _srt(tr: Transcript, path: Path) -> None:
    out = []
    for i, (s, e, spk, text) in enumerate(_cues(tr), 1):
        label = f"{tr.speaker_label(spk)}: " if tr.diarized else ""
        out.append(f"{i}\n{fmt_time(s, True, ',')} --> {fmt_time(e, True, ',')}\n{_wrap(label + text)}\n")
    path.write_text("\n".join(out), encoding="utf-8")


def _vtt(tr: Transcript, path: Path) -> None:
    out = ["WEBVTT", ""]
    for s, e, spk, text in _cues(tr):
        voice = f"<v {tr.speaker_label(spk)}>" if tr.diarized else ""
        out.append(f"{fmt_time(s, True)} --> {fmt_time(e, True)}\n{voice}{_wrap(text)}\n")
    path.write_text("\n".join(out), encoding="utf-8")


def _json(tr: Transcript, path: Path) -> None:
    data = {
        "arquivo": Path(tr.source).name,
        "idioma": tr.language,
        "duracao_segundos": round(tr.duration, 3),
        "modelo": tr.model,
        "falantes": {str(k): v for k, v in tr.speaker_names.items()},
        "falas": [
            {
                "inicio": round(u.start, 3),
                "fim": round(u.end, 3),
                "falante": tr.speaker_label(u.speaker) or None,
                "texto": u.text,
                "palavras": [
                    {"inicio": round(w.start, 3), "fim": round(w.end, 3), "texto": w.text.strip()}
                    for w in u.words
                ],
            }
            for u in tr.utterances
        ],
    }
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _docx(tr: Transcript, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(Path(tr.source).stem, level=1)
    meta = doc.add_paragraph()
    info = [
        f"Duração: {fmt_time(tr.duration)}",
        f"Idioma: {LANG_NAMES.get(tr.language, tr.language)}",
    ]
    if tr.diarized:
        info.append(f"Falantes: {', '.join(tr.speaker_names.values())}")
    info.append(f"Gerado em {datetime.now():%d/%m/%Y %H:%M}")
    run = meta.add_run(" · ".join(info))
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    palette = [RGBColor(0x1F, 0x5F, 0xBF), RGBColor(0xB3, 0x40, 0x1A), RGBColor(0x2E, 0x7D, 0x32),
               RGBColor(0x7B, 0x1F, 0xA2), RGBColor(0x00, 0x83, 0x8F), RGBColor(0xAD, 0x14, 0x57)]

    for u in tr.utterances:
        p = doc.add_paragraph()
        ts = p.add_run(f"[{fmt_time(u.start)}] ")
        ts.font.size = Pt(9)
        ts.font.color.rgb = RGBColor(0x90, 0x90, 0x90)
        if tr.diarized:
            name = p.add_run(f"{tr.speaker_label(u.speaker)}: ")
            name.bold = True
            if u.speaker is not None:
                name.font.color.rgb = palette[u.speaker % len(palette)]
        p.add_run(u.text)
    doc.save(str(path))
